import csv
from io import BytesIO, StringIO
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pypdf import PdfReader


def extract_text(file_bytes: bytes, filename: str):
    extension = Path(filename).suffix.lower()
    sections = []

    if extension == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))

        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()

            if text:
                sections.append({
                    "text": text,
                    "metadata": {"page": page_number},
                })

    elif extension == ".docx":
        document = Document(BytesIO(file_bytes))
        parts = []

        # Read paragraphs and tables in document order.
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                text = block.text.strip()

                if text:
                    parts.append(text)
            else:
                # The block is a table.
                for row in block.rows:
                    cells = [cell.text.strip() for cell in row.cells]

                    if any(cells):
                        parts.append(" | ".join(cells))

        text = "\n\n".join(parts)

        if text:
            sections.append({
                "text": text,
                "metadata": {"sourceType": "docx"},
            })

    elif extension == ".txt":
        text = file_bytes.decode("utf-8-sig").strip()

        if text:
            sections.append({
                "text": text,
                "metadata": {"sourceType": "txt"},
            })

    elif extension == ".csv":
        text = file_bytes.decode("utf-8-sig")
        reader = csv.reader(StringIO(text, newline=""))

        # This version expects the first row to contain column headers.
        headers = next(reader, [])

        for row_number, row in enumerate(reader, start=2):
            if not any(value.strip() for value in row):
                continue

            fields = []

            for index, value in enumerate(row):
                header = (
                    headers[index].strip()
                    if index < len(headers)
                    else ""
                )
                label = header or f"Column {index + 1}"
                fields.append(f"{label}: {value}")

            sections.append({
                "text": "\n".join(fields),
                "metadata": {"row": row_number},
            })

    elif extension == ".xlsx":
        workbook = load_workbook(
            BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )

        try:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)

                # This version expects headers in the first row.
                headers = next(rows, ())

                for row_number, row in enumerate(rows, start=2):
                    if all(
                        value is None or str(value).strip() == ""
                        for value in row
                    ):
                        continue

                    fields = []

                    for index, value in enumerate(row):
                        header = (
                            headers[index]
                            if index < len(headers)
                            else None
                        )
                        label = (
                            str(header).strip()
                            if header is not None
                            else ""
                        )
                        label = label or f"Column {index + 1}"
                        value_text = "" if value is None else str(value)

                        fields.append(f"{label}: {value_text}")

                    sections.append({
                        "text": "\n".join(fields),
                        "metadata": {
                            "sheet": sheet.title,
                            "row": row_number,
                        },
                    })
        finally:
            workbook.close()

    else:
        raise ValueError(f"Unsupported file type: {extension}")

    if not sections:
        raise ValueError(f"No readable text found in {filename}")

    return sections